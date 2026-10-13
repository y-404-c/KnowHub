from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


out = Path(__file__).resolve().parents[1] / "data" / "samples" / "面试演示产品方案.docx"
out.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
run = title.add_run("知枢 KnowHub 智能知识库产品方案")
run.bold = True
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(46, 116, 181)

subtitle = doc.add_paragraph("面试演示用产品资料｜2026 年 8 月")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)

for heading, body in [
    ("一、产品定位", "知枢 KnowHub 是面向企业内部知识管理的 RAG 平台。它支持 Word、PDF、Markdown、TXT 与图片 OCR 上传，并通过混合检索、重排序和引用溯源回答问题。"),
    ("二、核心能力", "1. 文档解析与知识入库：上传后自动清洗、切分并生成向量索引。\n2. 混合检索：结合 BM25 关键词检索与向量语义检索，再用 RRF 融合结果。\n3. 多 Agent：Planner 识别检索、计算和日期工具需求；Retriever 实际执行工具；Reviewer 审查结果。\n4. 流式回答：以 SSE 实时返回答案、引用来源和 Agent 执行轨迹。"),
    ("三、商务规则", "企业版年费为 1280 元。年度促销活动按原价的 85% 收费；如客户购买 20 个账号，则总价等于折后单价乘以账号数量。"),
    ("四、实施建议", "建议先导入制度、产品手册与常见问答，再通过带引用的问答验证检索质量。涉及价格计算时，必须调用 calculator 工具，避免由语言模型直接心算。"),
]:
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(heading)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(46, 116, 181)
    p = doc.add_paragraph(body)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15

doc.save(out)
print(out)
